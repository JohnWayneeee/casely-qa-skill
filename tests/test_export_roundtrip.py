"""Round-trip check for the XLSX and CSV exports.

Matching headers alone does not catch damaged text or a step split into extra
rows. Each test exports the fixtures, reads the file back and compares every
cell with the text parsed from the untouched Markdown fixtures.

Run from the repository root: python -m unittest discover tests
"""

import csv
import hashlib
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "skills" / "casely" / "scripts"))

import export_to_xlsx as exporter  # noqa: E402

FIXTURES = Path(__file__).resolve().parent / "fixtures" / "multilingual"


def _digest(folder: Path) -> dict:
    return {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in sorted(folder.glob("*.md"))}


class ExportRoundTrip(unittest.TestCase):
    def setUp(self) -> None:
        # Export from a copy so the fixtures stay the reference for comparison.
        self.tmp = Path(tempfile.mkdtemp())
        self.results = self.tmp / "results"
        shutil.copytree(FIXTURES, self.results)
        self.out = self.tmp / "exports"
        self.out.mkdir()
        self.fixture_digest = _digest(FIXTURES)

        self.cases, problems = exporter.read_cases(self.results)
        self.assertEqual(problems, [])
        self.headers = self.cases[0][1]
        self.expected = [
            [exporter.cell_text(value) for value in row]
            for _, _, rows in self.cases
            for row in rows
        ]

    def tearDown(self) -> None:
        shutil.rmtree(self.tmp)
        self.assertEqual(_digest(FIXTURES), self.fixture_digest, "the export changed a fixture")

    def test_expected_text_keeps_the_hard_parts(self) -> None:
        flat = "\n".join("\t".join(row) for row in self.expected)
        for needle in ("메뉴판 주세요, 부탁해요", '"김치찌개, 2인분"', "A | B", "50,000 ريال", '"خاطئة، مرتين"'):
            self.assertIn(needle, flat)
        self.assertTrue(all("<br>" not in cell for row in self.expected for cell in row))
        self.assertTrue(any("\n" in cell for row in self.expected for cell in row))

    def test_xlsx_round_trip(self) -> None:
        dest = exporter.export_combined(self.cases, self.out, "all_test_cases.xlsx")
        ws = load_workbook(dest).active
        rows = [[cell if cell is not None else "" for cell in row] for row in ws.iter_rows(values_only=True)]

        self.assertEqual(rows[0], self.headers)
        self.assertEqual(len(rows) - 1, len(self.expected), "row count changed")
        self.assertEqual(ws.max_column, len(self.headers), "column count changed")
        self.assertEqual(rows[1:], self.expected)

    def test_csv_round_trip(self) -> None:
        dest = exporter.export_csv(self.cases, self.out, "all_test_cases.csv")
        self.assertTrue(dest.read_bytes().startswith(b"\xef\xbb\xbf"), "CSV lost its UTF-8 BOM")

        with dest.open(encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.reader(handle))

        self.assertEqual(rows[0], self.headers)
        self.assertEqual(len(rows) - 1, len(self.expected), "row count changed")
        self.assertTrue(all(len(row) == len(self.headers) for row in rows), "column count changed")
        self.assertEqual(rows[1:], self.expected)

    def test_csv_refuses_mixed_columns(self) -> None:
        (self.results / "ZZ-001.md").write_text(
            "| ID | Title |\n|----|-------|\n| ZZ-001 | 다른 열 |\n", encoding="utf-8"
        )
        cases, _ = exporter.read_cases(self.results)
        with self.assertRaises(exporter.TableError):
            exporter.export_csv(cases, self.out, "all_test_cases.csv")


if __name__ == "__main__":
    unittest.main()
